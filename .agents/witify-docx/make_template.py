#!/usr/bin/env python3
"""Maintainer-only: regenerate skills/productivity/witify-docx/template/witify-template.docx.

The template carries the whole brand (styles, cover, header, footer). Day-to-day
documents never touch this script: the skill opens the template with python-docx
and only writes content. Re-run this after changing the assets or the layout
here; alternatively, edit the template directly in Word and leave this script
out of date (note it in the commit).

Requires: python-docx, pymupdf, numpy, Pillow.
"""
import os
import sys

import numpy as np
import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageEnhance, ImageFilter

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
SKILL_TEMPLATE_DIR = os.path.join(HERE, "..", "..", "skills", "productivity", "witify-docx", "template")
OUT = os.path.abspath(os.path.join(SKILL_TEMPLATE_DIR, "witify-template.docx"))

FONT, MONO = "Helvetica", "Consolas"
HEADING_FONT = "Archivo"
RED = (1, 0.278, 0.275)                                          # FF4746
GRAY = RGBColor(0x5E, 0x6A, 0x6E)
LIGHT = RGBColor(0xAE, 0xBD, 0xC2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
EMU_PER_INCH = 914400
EMU_PER_PT = 12700
BAND_COLOR, BAND_HEIGHT_PT = "1D373D", 4.5                       # teal band on content pages


# ============================================================= cover background (one image, no text)
COVER_PX = (1700, 2200)                                                  # 850x1100 pt at dpi 144
COVER_BASE, COVER_DARK, COVER_TEAL = (0x1D, 0x37, 0x3D), (0x0B, 0x2A, 0x31), (0x1C, 0x5A, 0x6A)
GRAIN_SIGMA = 4.0                                                        # luminance grain, on 0-255


def render_cover_gradient():
    """Soft teal field around COVER_BASE: darkest top-right, lighter teal bottom-left and
    bottom-right, one diagonal blend, then a fine uniform luminance grain."""
    w, h = COVER_PX
    x = np.linspace(0.0, 1.0, w, dtype=np.float32)[None, :]
    y = np.linspace(0.0, 1.0, h, dtype=np.float32)[:, None]

    def blob(cx, cy, sx, sy):
        return np.exp(-(((x - cx) / sx) ** 2 + ((y - cy) / sy) ** 2))

    dark = 0.80 * blob(0.75, 0.15, 0.50, 0.40)                           # top-right shadow
    teal = 0.70 * blob(-0.05, 0.80, 0.30, 0.45) + 0.60 * blob(1.05, 1.05, 0.30, 0.35)  # bottom corners
    teal = np.clip(teal + 0.20 * np.clip((y - x + 0.2), 0, 1), 0, 1)      # gentle diagonal lift
    base = np.array(COVER_BASE, np.float32)
    img = (base + dark[..., None] * (np.array(COVER_DARK, np.float32) - base)
                + teal[..., None] * (np.array(COVER_TEAL, np.float32) - base))
    rng = np.random.default_rng(1103)
    img += rng.normal(0.0, GRAIN_SIGMA, (h, w, 1)).astype(np.float32)
    return Image.fromarray(np.clip(img, 0, 255).astype(np.uint8))


def svg_alpha(name, size_px):
    """Rasterize an SVG to an RGBA Pillow image at the given pixel size."""
    src = pymupdf.open(os.path.join(ASSETS, name))
    zoom = (size_px[0] / src[0].rect.width, size_px[1] / src[0].rect.height)
    pix = src[0].get_pixmap(matrix=pymupdf.Matrix(*zoom), alpha=True)
    return Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)


def render_cover_background(path):
    """Grained teal gradient, light logo top-left, red accent, tone-on-tone W watermark bleeding
    bottom-right. The watermark is the local background lifted one tone, at partial opacity, so it
    follows the gradient instead of fighting it."""
    ground = render_cover_gradient()

    # Watermark: same rect as before (pt 300..1060 x 722.5..1220), scaled 2x to pixels.
    wm_box = (600, 1445, 2120, 2440)
    mark = svg_alpha("witify-icon-tonal.svg", (wm_box[2] - wm_box[0], wm_box[3] - wm_box[1]))
    mark = mark.crop((0, 0, min(mark.width, ground.width - wm_box[0]), min(mark.height, ground.height - wm_box[1])))
    region = ground.crop((wm_box[0], wm_box[1], wm_box[0] + mark.width, wm_box[1] + mark.height))
    lifted = ImageEnhance.Brightness(region).enhance(1.26)
    alpha = mark.getchannel("A").filter(ImageFilter.GaussianBlur(1.2)).point(lambda a: int(a * 0.45))
    ground.paste(Image.composite(lifted, region, alpha), (wm_box[0], wm_box[1]))

    ground_jpg = os.path.join(HERE, "cover-ground.jpg")
    ground.save(ground_jpg, quality=90, subsampling=0)

    pdf = pymupdf.open()
    page = pdf.new_page(width=850, height=1100)
    page.insert_image(page.rect, filename=ground_jpg)

    def svg(name, rect):
        src = pymupdf.open("pdf", pymupdf.open(os.path.join(ASSETS, name)).convert_to_pdf())
        page.show_pdf_page(rect, src, 0)

    svg("logo-light.svg", pymupdf.Rect(72, 64, 240, 106.3))
    page.draw_rect(pymupdf.Rect(72, 440, 136, 444), fill=RED, color=None)
    page.get_pixmap(dpi=144).save(path, jpg_quality=90)
    os.remove(ground_jpg)


# ============================================================= XML helpers
def set_fonts(rpr, name):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    # Theme font attributes beat explicit names in Word; strip them.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def style_font(style, name=FONT, size=None, bold=None, color="auto"):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    set_fonts(rpr, name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    existing = rpr.find(qn("w:color"))
    if existing is not None:
        rpr.remove(existing)
    if color == "auto":                                       # black text, overriding heading blues
        el = OxmlElement("w:color")
        el.set(qn("w:val"), "auto")
        rpr.append(el)
    else:
        style.font.color.rgb = color


def anchor_full_page(inline_shape, width_in, height_in):
    """Turn an inline picture into a floating one at the page's top-left, behind text."""
    inline = inline_shape._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
                       ("simplePos", "0"), ("relativeHeight", "0"), ("behindDoc", "1"),
                       ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1")):
        anchor.set(key, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for tag in ("wp:positionH", "wp:positionV"):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        pos.append(offset)
        anchor.append(pos)
    extent = inline.find(qn("wp:extent"))
    extent.set("cx", str(int(width_in * EMU_PER_INCH)))
    extent.set("cy", str(int(height_in * EMU_PER_INCH)))
    anchor.append(extent)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(inline.find(qn("wp:docPr")))
    frame_pr = inline.find(qn("wp:cNvGraphicFramePr"))
    if frame_pr is not None:
        anchor.append(frame_pr)
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)


def add_top_band(paragraph):
    """Teal band flush with the top edge of the page: a full-width rectangle anchored to the page
    from the header, behind the text. A section page border (w:pgBorders, offsetFrom="page") was
    the obvious tool, but LibreOffice lays that box out around the text area, so its top line
    stops at the side margins in every PDF it exports; an anchored shape renders edge to edge."""
    xml = f"""
    <wp:anchor {nsdecls("wp", "a")} xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
               distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1"
               locked="0" layoutInCell="1" allowOverlap="1">
      <wp:simplePos x="0" y="0"/>
      <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
      <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
      <wp:extent cx="{int(8.5 * EMU_PER_INCH)}" cy="{int(BAND_HEIGHT_PT * EMU_PER_PT)}"/>
      <wp:effectExtent l="0" t="0" r="0" b="0"/>
      <wp:wrapNone/>
      <wp:docPr id="2" name="Top band"/>
      <wp:cNvGraphicFramePr/>
      <a:graphic>
        <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
          <wps:wsp>
            <wps:cNvSpPr/>
            <wps:spPr>
              <a:xfrm><a:off x="0" y="0"/><a:ext cx="{int(8.5 * EMU_PER_INCH)}" cy="{int(BAND_HEIGHT_PT * EMU_PER_PT)}"/></a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
              <a:solidFill><a:srgbClr val="{BAND_COLOR}"/></a:solidFill>
              <a:ln><a:noFill/></a:ln>
            </wps:spPr>
            <wps:bodyPr/>
          </wps:wsp>
        </a:graphicData>
      </a:graphic>
    </wp:anchor>"""
    drawing = OxmlElement("w:drawing")
    drawing.append(parse_xml(xml))
    paragraph.add_run()._r.append(drawing)


def add_page_field(paragraph):
    run = paragraph.add_run()
    for el_name, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "1"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        el = OxmlElement(el_name)
        for attr, value in attrs.items():
            el.set(qn(attr), value)
        if text is not None:
            el.text = text
        run._element.append(el)
    return run


def pin_to_page(paragraph, y_in, width_twips=10080):
    """Wrap the paragraph in a Word frame anchored at y_in from the top of the page, left margin,
    text-width wide. Same-framePr neighbours share the frame."""
    ppr = paragraph._p.get_or_add_pPr()
    frame = OxmlElement("w:framePr")
    for attr, value in (("w:w", str(width_twips)), ("w:hRule", "auto"), ("w:wrap", "notBeside"),
                        ("w:vAnchor", "page"), ("w:hAnchor", "margin"), ("w:xAlign", "left"),
                        ("w:y", str(int(y_in * 1440)))):
        frame.set(qn(attr), value)
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is not None:
        pstyle.addnext(frame)
    else:
        ppr.insert(0, frame)


def set_compatibility_mode(doc, val):
    """Set (or add) w:compat/w:compatSetting[compatibilityMode] in settings.xml."""
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    for setting in compat.findall(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), val)
            return
    setting = OxmlElement("w:compatSetting")
    setting.set(qn("w:name"), "compatibilityMode")
    setting.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    setting.set(qn("w:val"), val)
    compat.append(setting)


def paragraph_shading(style, fill_hex):
    ppr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill_hex)
    ppr.append(shd)


def el(tag, attrs=None, *children):
    """OxmlElement with w:-prefixed attributes and children."""
    node = OxmlElement(tag)
    for attr, value in (attrs or {}).items():
        node.set(qn(f"w:{attr}"), str(value))
    for child in children:
        node.append(child)
    return node


def add_table_style(doc):
    """'Witify Table': horizontal rules only, bold header on light gray with a dark bottom rule.
    Word applies the firstRow part only when the table's w:tblLook has firstRow="1"."""
    def border(tag, **attrs):
        return el(f"w:{tag}", attrs)
    thin = dict(val="single", sz="4", space="0", color="D0D5D3")
    doc.styles.element.append(el("w:style", {"type": "table", "styleId": "WitifyTable"},
        el("w:name", {"val": "Witify Table"}),
        el("w:basedOn", {"val": "TableNormal"}),
        el("w:uiPriority", {"val": "59"}),
        el("w:pPr", None, el("w:spacing", {"before": "40", "after": "40", "line": "240", "lineRule": "auto"})),
        el("w:rPr", None, el("w:rFonts", {"ascii": FONT, "hAnsi": FONT, "cs": FONT}), el("w:sz", {"val": "20"})),
        el("w:tblPr", None,
           el("w:tblBorders", None, border("top", **thin), border("bottom", **thin), border("insideH", **thin),
              border("left", val="nil"), border("right", val="nil"), border("insideV", val="nil")),
           el("w:tblCellMar", None, *(el(f"w:{side}", {"w": w, "type": "dxa"})
                                      for side, w in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100))))),
        el("w:tblStylePr", {"type": "firstRow"},
           el("w:rPr", None, el("w:b")),
           el("w:tcPr", None,
              el("w:tcBorders", None, border("bottom", val="single", sz="8", space="0", color="1D373D")),
              el("w:shd", {"val": "clear", "color": "auto", "fill": "F3F4F2"})))))


# ============================================================= build
def main():
    os.makedirs(SKILL_TEMPLATE_DIR, exist_ok=True)
    cover_jpg = os.path.join(HERE, "cover-background.jpg")
    icon_png = os.path.join(ASSETS, "witify-icon-black50.png")
    render_cover_background(cover_jpg)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("top", "bottom", "left", "right"):
        setattr(section, f"{side}_margin", Inches(0.75))
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    section.different_first_page_header_footer = True

    # Content pages number from 1: the cover is page 0 and shows no number.
    # sectPr schema order matters to Word: pgMar, pgNumType, then cols/titlePg/docGrid.
    section._sectPr.find(qn("w:pgMar")).addnext(el("w:pgNumType", {"start": "0"}))

    # ---- styles
    styles = doc.styles
    style_font(styles["Normal"], size=11)
    style_font(styles["Title"], name=HEADING_FONT, size=24, bold=True)
    style_font(styles["Heading 1"], name=HEADING_FONT, size=16, bold=True)
    style_font(styles["Heading 2"], name=HEADING_FONT, size=13, bold=True)
    style_font(styles["Heading 3"], name=HEADING_FONT, size=11.5, bold=True)
    style_font(styles["List Bullet"], size=11)
    style_font(styles["List Number"], size=11)
    for name, before, after in (("Heading 1", 29, 12), ("Heading 2", 12, 7), ("Heading 3", 12, 0)):
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
    # Air between list items. The stock List Bullet style carries contextualSpacing, which
    # suppresses the space between consecutive items of the same style; drop it.
    for name in ("List Bullet", "List Number"):
        styles[name].paragraph_format.space_after = Pt(4)
        ppr = styles[name].element.get_or_add_pPr()
        for contextual in ppr.findall(qn("w:contextualSpacing")):
            ppr.remove(contextual)

    def new_par_style(style_name, base="Normal", **font):
        st = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles[base]
        st.quick_style = True
        style_font(st, **font)
        return st

    cover_version = new_par_style("Cover Version", size=11, bold=False, color=LIGHT)
    cover_version.paragraph_format.space_before = Inches(4.05) - Pt(1)   # the 1pt background paragraph sits above
    cover_version.paragraph_format.space_after = Pt(0)
    cover_version.paragraph_format.keep_with_next = True

    cover_title = new_par_style("Cover Title", name=HEADING_FONT, size=42, bold=True, color=WHITE)
    cover_title.paragraph_format.space_before = Pt(4)
    cover_title.paragraph_format.space_after = Pt(0)
    cover_title.paragraph_format.line_spacing = 1.0
    cover_title.paragraph_format.keep_with_next = True

    cover_subtitle = new_par_style("Cover Subtitle", size=14.5, color=LIGHT)
    cover_subtitle.paragraph_format.space_before = Pt(14)
    cover_subtitle.paragraph_format.space_after = Pt(0)

    cover_meta = new_par_style("Cover Meta", size=9.5, bold=False, color=WHITE)
    cover_meta.paragraph_format.space_after = Pt(4)

    code = new_par_style("Code", name=MONO, size=9)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Pt(6)
    code.paragraph_format.keep_with_next = True
    paragraph_shading(code, "F3F4F2")

    table_cell = new_par_style("Table Cell", size=10)
    table_cell.paragraph_format.space_before = Pt(2)
    table_cell.paragraph_format.space_after = Pt(2)
    table_sub = new_par_style("Table Sub", size=9.5, color=GRAY)
    table_sub.paragraph_format.space_before = Pt(0)
    table_sub.paragraph_format.space_after = Pt(2)

    code_inline = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
    style_font(code_inline, name=MONO, size=10)
    add_table_style(doc)

    # ---- cover: background image anchored in the body (Word for Mac dims header content while
    # the body is edited, so a header-anchored cover looks washed out on screen), live text below.
    # The picture sits in its own 1pt paragraph so build_docx.py still finds the bare "{{title}}"
    # paragraph; the first-page header stays empty (the different-first-page setting still gates
    # the icon header off the cover).
    bg = doc.add_paragraph()
    bg.paragraph_format.space_before = Pt(0)
    bg.paragraph_format.space_after = Pt(0)
    bg.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    bg.paragraph_format.line_spacing = Pt(1)
    bg_run = bg.add_run()
    bg_run.font.size = Pt(1)
    shape = bg_run.add_picture(cover_jpg, width=Inches(8.5))
    anchor_full_page(shape, 8.5, 11)

    # Explicit (empty) first-page header and footer parts, so the cover shows neither the icon
    # header nor the page-number footer.
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False

    doc.add_paragraph("{{version}}", style=cover_version)
    doc.add_paragraph("{{title}}", style=cover_title)
    doc.add_paragraph("{{subtitle}}", style=cover_subtitle)

    # Prepared-for / prepared-by block: a white rule then two columns (tab stop at 3.5in), body
    # paragraphs pinned to the bottom of the cover with a Word paragraph frame (consecutive
    # paragraphs with identical framePr share one frame). It lives in the body, not the first-page
    # footer, because LibreOffice paints footers before body-anchored behindDoc pictures, which
    # would hide the text under the cover background in its render. One run per text piece so
    # build_docx.py can replace placeholders inside runs.
    rule = doc.add_paragraph(style=cover_meta)
    rule.paragraph_format.space_after = Pt(10)
    block = [rule]
    for left, right_, bold in (("Préparé pour", "Préparé par", True), ("{{client}}", "{{author}}", False),
                               ("{{date}}", "Witify", False)):
        p = doc.add_paragraph(style=cover_meta)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), "5040")
        tabs = OxmlElement("w:tabs")
        tabs.append(tab)
        p._p.get_or_add_pPr().append(tabs)
        for text in (left, "\t", right_):
            p.add_run(text).bold = bold
        block.append(p)
    # The white rule is a top border on the empty first paragraph. Word and LibreOffice merge
    # identical borders on adjacent paragraphs into one box, so every paragraph of the block
    # carries it and one line is drawn; LibreOffice drops the border when only the first has it.
    for p in block:
        pin_to_page(p, y_in=9.35)
        top = OxmlElement("w:top")
        for attr, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "FFFFFF")):
            top.set(qn(attr), value)
        border = OxmlElement("w:pBdr")
        border.append(top)
        p._p.pPr.find(qn("w:framePr")).addnext(border)

    doc.add_page_break()

    # ---- content pages: icon top-right, brand + page number in the footer
    icon_p = section.header.paragraphs[0]
    icon_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    icon_p.paragraph_format.space_after = Pt(0)
    add_top_band(icon_p)
    icon_p.add_run().add_picture(icon_png, height=Inches(0.2))

    footer_p = section.footer.paragraphs[0]
    tabs = OxmlElement("w:tabs")
    for pos in ("4680", "9360"):
        clear = OxmlElement("w:tab")
        clear.set(qn("w:val"), "clear")
        clear.set(qn("w:pos"), pos)
        tabs.append(clear)
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:pos"), "10080")
    tabs.append(right)
    footer_p.paragraph_format.element.get_or_add_pPr().append(tabs)
    for run in (footer_p.add_run("Witify Technologies · witify.io"), footer_p.add_run("\t"), add_page_field(footer_p)):
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY

    # ---- settings: declare Word 2013+ compatibility so Word doesn't show "Compatibility Mode"
    set_compatibility_mode(doc, "15")

    doc.save(OUT)
    os.remove(cover_jpg)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    sys.exit(main())
